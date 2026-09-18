"""End-to-End Pipeline Verification across FSF-M1, M2, M3, M4."""

from __future__ import annotations

import json
from pathlib import Path
from docx import Document

from fsf.cli.main import main
from fsf.conformance.media import (
    build_expected_contract,
    evaluate_media_file,
    map_conformance_status_to_host,
)
from fsf.executors.registry import get_executor
from fsf.intake.compiler import compile_job
from fsf.runtime.external_op import build_resumed_job, inspect_existing_artifacts
from fsf.runtime.plan import build_execution_plan


class E2EMockProbe:
    """Mock probe returning conforming video facts matching 1024x576, h264, 24fps."""

    def probe(self, path: Path) -> dict:
        return {
            "format": {"duration": "5.0", "format_name": "mov,mp4"},
            "streams": [
                {
                    "index": 0,
                    "codec_type": "video",
                    "codec_name": "h264",
                    "width": 1024,
                    "height": 576,
                    "avg_frame_rate": "24/1",
                    "nb_frames": "120",
                }
            ],
            "black_frame_ratio": 0.0,
            "freeze_ranges": [],
        }


def test_full_lifecycle_e2e(tmp_path: Path) -> None:
    """Full pipeline verification:
    1. DOCX Intake (M1)
    2. Execution Plan Compilation (M2)
    3. Multi-Executor Generation (M4)
    4. Durable Checkpoint & Partial Resume (M2)
    5. Media Conformance Consumption (M3)
    """
    # Create test DOCX
    docx_path = tmp_path / "cinematic_treatment.docx"
    doc = Document()
    doc.add_paragraph("Dramatic drone establishing shot of misty pine forest at sunrise. 5.0 seconds")
    doc.add_paragraph("-----")
    doc.add_paragraph("Close-up macro shot of dew drop falling from a green leaf. 5 seconds")
    doc.save(docx_path)

    # --- Phase 1: M1 Intake ---
    out_dir = tmp_path / "artifacts"
    manifest, request, profile = compile_job(docx_path, out_dir=out_dir)

    assert manifest["schema_version"] == "fsf_video_job_v1"
    assert manifest["task_count"] == 2
    assert profile["schema_version"] == "prodocux_docx_profile_v1"
    assert request["schema_version"] == "pdx_tool_request_v1"

    # --- Phase 2: M2 Execution Plan ---
    plan = build_execution_plan(manifest, executor_tool="studio.video.generate")
    assert plan["schema_version"] == "pdx_execution_plan_v1"
    assert len(plan["steps"]) == 4  # 2 tasks * 2 steps

    # --- Phase 3: M4 Generation & Checkpointing ---
    executor = get_executor("mock.studio.video")

    # Step A: Execute only shot 1 first
    res_1 = executor.execute_shot(manifest["tasks"][0], out_dir)
    assert res_1["schema_version"] == "pdx_tool_result_v1"
    assert (out_dir / "task_001.mock.bin").is_file()

    # Step B: Resume check — task 1 should be reused, task 2 pending
    resumed_manifest, pending, completed = build_resumed_job(manifest, out_dir)
    assert len(completed) == 1
    assert completed[0]["id"] == "task_001"
    assert len(pending) == 1
    assert pending[0]["id"] == "task_002"

    # Step C: Execute remaining pending task 2
    res_2 = executor.execute_shot(pending[0], out_dir)
    assert res_2["schema_version"] == "pdx_tool_result_v1"
    assert (out_dir / "task_002.mock.bin").is_file()

    # Verify both artifacts now exist
    status_final = inspect_existing_artifacts(manifest, out_dir)
    assert status_final["task_001"]["completed"] is True
    assert status_final["task_002"]["completed"] is True

    # --- Phase 4: M3 Media Conformance ---
    for task in manifest["tasks"]:
        shot_id = task["id"]
        shot_file = out_dir / f"{shot_id}.mock.bin"
        expected = build_expected_contract(task, width=1024, height=576, fps=24)

        report, host_action = evaluate_media_file(
            shot_file,
            expected,
            probe_runner=E2EMockProbe(),
            request_id=f"req:qc:{shot_id}",
        )

        assert report["evaluation_status"] == "conforms"
        assert host_action["host_action"] == "permit_next_step"
        assert host_action["create_binding"] is True

    # --- Phase 5: CLI E2E Verification ---
    cli_out = tmp_path / "cli_pipeline"
    rc = main(["run", str(docx_path), "--out", str(cli_out), "--executor", "mock.studio.video"])
    assert rc == 0
    assert (cli_out / "task_001.mock.bin").is_file()
    assert (cli_out / "task_002.mock.bin").is_file()
    assert (cli_out / "fsf_checkpoint.json").is_file()
