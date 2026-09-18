"""Edge case tests for FSF-M1 intake compiler."""

from __future__ import annotations

import json
from pathlib import Path
import pytest
from docx import Document

from fsf.intake.compiler import (
    compile_job,
    compile_job_from_profile,
    compile_tasks_from_profile,
)


def test_custom_durations_and_whitespace(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("   Shot with leading space and 12 seconds.   ")
    doc.add_paragraph("   ")  # empty whitespace paragraph
    doc.add_paragraph("-----")
    doc.add_paragraph("-----")  # consecutive separators
    doc.add_paragraph("Second shot with no duration specified.")
    doc.add_paragraph("-----")
    doc.add_paragraph("Third shot with short duration 2.5s.")

    docx_path = tmp_path / "custom.docx"
    doc.save(docx_path)

    manifest, request, profile = compile_job(docx_path, tool_name="custom.studio.video")
    assert manifest["task_count"] == 3
    assert request["tool"] == "custom.studio.video"

    t1, t2, t3 = manifest["tasks"]
    assert t1["duration_seconds"] == 12.0
    assert t1["normalized_prompt"] == "Shot with leading space and"

    assert t2["duration_seconds"] == 5.0  # default
    assert t2["normalized_prompt"] == "Second shot with no duration specified."

    assert t3["duration_seconds"] == 2.5
    assert t3["normalized_prompt"] == "Third shot with short duration"


def test_direct_synthetic_profile_consumption() -> None:
    """Verifies that the compiler can run purely on a mock/synthetic profile dict with zero docx file."""
    mock_profile = {
        "schema_version": "prodocux_docx_profile_v1",
        "source": {
            "name": "virtual_script.docx",
            "sha256": "abcdef1234567890abcdef1234567890abcdef1234567890abcdef1234567890",
        },
        "paragraph_count": 3,
        "heading_count": 0,
        "table_count": 0,
        "section_count": 1,
        "paragraphs": [
            {"index": 1, "style": "Normal", "text": "First scene prompt 7 seconds"},
            {"index": 2, "style": "Normal", "text": "-----"},
            {"index": 3, "style": "Normal", "text": "Second scene prompt"},
        ],
        "tables": [],
        "headers": [],
        "footers": [],
        "preview_truncated": False,
        "interpretation": "none",
    }

    manifest = compile_job_from_profile(mock_profile)
    assert manifest["source_document"] == "virtual_script.docx"
    assert manifest["source_sha256"] == "abcdef1234567890abcdef1234567890abcdef1234567890abcdef1234567890"
    assert manifest["task_count"] == 2
    assert manifest["tasks"][0]["duration_seconds"] == 7.0
    assert manifest["tasks"][1]["duration_seconds"] == 5.0
    assert manifest["tasks"][0]["shot_spec"]["provenance"]["source_sha256"] == mock_profile["source"]["sha256"]


def test_invalid_profile_schema() -> None:
    bad_profile = {
        "schema_version": "unsupported_schema_v9",
        "paragraphs": [],
    }
    with pytest.raises(ValueError, match="Unsupported profile schema"):
        compile_tasks_from_profile(bad_profile)
