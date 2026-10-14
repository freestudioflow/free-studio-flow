"""FSF-owned shot/asset labels and compiled expected media contracts."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from fsf.cli.main import main
from fsf.conformance.media import canonical_contract_digest, expected_contract_for_task
from fsf.intake.compiler import compile_job, compile_job_from_profile
from fsf.intake.parser import classify_unit, skill_applies_to_task
from fsf.runtime.external_op import build_resumed_job, inspect_existing_artifacts
from fsf.runtime.plan import build_execution_plan
from fsf.studio.runner import run_studio_flow, select_tasks_for_skill
from fsf.runtime.plan import build_execution_plan
from fsf.studio.runner import run_studio_flow, select_tasks_for_skill


def _mixed_sheet(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("SHOT: Locked-off wide shot of mist over water. 5 seconds")
    doc.add_paragraph("-----")
    doc.add_paragraph("ASSET: hero vehicle mesh, no people")
    doc.add_paragraph("-----")
    doc.add_paragraph("Unlabeled ocean plate. 8 seconds")
    path = tmp_path / "mixed_units.docx"
    doc.save(path)
    return path


def test_classify_requires_colon_on_body_text() -> None:
    kind, remainder, source = classify_unit("Shot with leading space and 12 seconds.")
    assert kind == "shot"
    assert source == "default_shot"
    assert remainder.startswith("Shot with leading space")

    kind, remainder, source = classify_unit("ASSET hero mesh without colon")
    assert kind == "shot"
    assert source == "default_shot"

    kind, remainder, source = classify_unit("ASSET: hero vehicle mesh")
    assert kind == "asset"
    assert source == "explicit_label"
    assert remainder == "hero vehicle mesh"

    kind, remainder, source = classify_unit("SHOT: misty forest 5 seconds")
    assert kind == "shot"
    assert source == "explicit_label"
    assert remainder.startswith("misty forest")


def test_classify_heading_may_omit_colon() -> None:
    kind, remainder, source = classify_unit("ASSET hero mesh", style="Heading 1")
    assert kind == "asset"
    assert source == "explicit_label"
    assert remainder == "hero mesh"


def test_compile_mixed_shot_and_asset(tmp_path: Path) -> None:
    manifest, request, _ = compile_job(_mixed_sheet(tmp_path))
    assert manifest["task_count"] == 3
    assert manifest["shot_count"] == 2
    assert manifest["asset_count"] == 1
    assert len(request["inputs"]["shots"]) == 2

    shot_a, asset, shot_b = manifest["tasks"]
    assert shot_a["unit_kind"] == "shot"
    assert shot_a["classification"] == "explicit_label"
    assert shot_a["mode"] == "text_to_video"
    assert "SHOT:" not in shot_a["prompt"]
    assert shot_a["expected_media_contract"]["duration_seconds"] == 5.0
    assert shot_a["shot_spec"]["schema_version"] == "fsf_shot_spec_v1"

    assert asset["unit_kind"] == "asset"
    assert asset["classification"] == "explicit_label"
    assert asset["mode"] == "asset_declare"
    assert "expected_media_contract" not in asset
    assert "shot_spec" not in asset
    assert asset["asset_spec"]["schema_version"] == "fsf_asset_spec_v1"
    assert asset["asset_spec"]["kind"] == "unspecified"
    assert asset["asset_spec"]["source_prompt"].startswith("hero vehicle")

    assert shot_b["unit_kind"] == "shot"
    assert shot_b["classification"] == "default_shot"
    assert shot_b["expected_media_contract"]["duration_seconds"] == 8.0
    stored = dict(shot_a["expected_media_contract"])
    stored["width"] = 1920
    assert expected_contract_for_task({**shot_a, "expected_media_contract": stored})["width"] == 1920


def test_heading_asset_from_kernel_profile() -> None:
    profile = {
        "schema_version": "prodocux_docx_profile_v1",
        "source": {
            "name": "heading_units.docx",
            "sha256": "abcdef1234567890abcdef1234567890abcdef1234567890abcdef1234567890",
        },
        "paragraph_count": 3,
        "heading_count": 1,
        "table_count": 0,
        "section_count": 1,
        "paragraphs": [
            {"index": 1, "style": "Heading 1", "text": "ASSET hero mesh"},
            {"index": 2, "style": "Normal", "text": "-----"},
            {"index": 3, "style": "Normal", "text": "SHOT: misty plate 5 seconds"},
        ],
        "tables": [],
        "headers": [],
        "footers": [],
        "preview_truncated": False,
        "interpretation": "none",
    }
    manifest = compile_job_from_profile(profile)
    assert manifest["shot_count"] == 1
    assert manifest["asset_count"] == 1
    assert manifest["tasks"][0]["unit_kind"] == "asset"
    assert manifest["tasks"][0]["classification"] == "explicit_label"
    assert manifest["tasks"][1]["unit_kind"] == "shot"
    assert manifest["tasks"][1]["expected_media_contract"]["duration_seconds"] == 5.0


def test_generate_and_p2_skip_asset_video_steps(tmp_path: Path) -> None:
    manifest, _, _ = compile_job(_mixed_sheet(tmp_path))
    generate = build_execution_plan(manifest, executor_tool="studio.video.generate")
    assert len(generate["steps"]) == 4
    assert all(
        step["inputs"].get("shot_id") != "task_002"
        for step in generate["steps"]
        if step["tool"] in {"studio.video.generate", "pdx.media.technical_conform"}
    )

    p2 = build_execution_plan(manifest, pipeline="p2")
    p2_by_task: dict[str, list[str]] = {}
    for step in p2["steps"]:
        shot_id = step["inputs"]["shot_id"]
        p2_by_task.setdefault(shot_id, []).append(step["tool"])
    assert "studio.video.generate" in p2_by_task["task_001"]
    assert "studio.video.generate" not in p2_by_task["task_002"]
    assert p2_by_task["task_002"] == ["studio.asset.3d", "studio.motion.mocap"]
    # 2 shots * (11 studio + 1 QC) + asset 3d/mocap
    assert len(p2["steps"]) == 26


def test_generate_flow_skips_asset_units(tmp_path: Path) -> None:
    manifest, _, _ = compile_job(_mixed_sheet(tmp_path))
    selected = select_tasks_for_skill(manifest["tasks"], "studio.video.generate")
    assert [task["id"] for task in selected] == ["task_001", "task_003"]
    assert skill_applies_to_task("studio.asset.3d", manifest["tasks"][1]) is True
    assert skill_applies_to_task("studio.video.generate", manifest["tasks"][1]) is False

    out = tmp_path / "gen"
    results = run_studio_flow(
        "studio.video.generate",
        manifest["tasks"],
        out,
        transport="local",
    )
    assert len(results) == 2
    names = {path.name for path in out.iterdir() if path.is_file()}
    assert any(name.startswith("task_001") for name in names)
    assert any(name.startswith("task_003") for name in names)
    assert not any(name.startswith("task_002") for name in names)


def test_cli_run_does_not_generate_assets(tmp_path: Path) -> None:
    sheet = _mixed_sheet(tmp_path)
    out = tmp_path / "run_out"
    assert main(["run", str(sheet), "--out", str(out), "--no-resume"]) == 0
    generated = {path.name for path in out.iterdir() if path.suffix == ".bin"}
    assert any(name.startswith("task_001") for name in generated)
    assert any(name.startswith("task_003") for name in generated)
    assert not any(name.startswith("task_002") for name in generated)


def test_plan_inputs_carry_mutually_exclusive_specs(tmp_path: Path) -> None:
    """P1-1: Plan steps must carry mutually exclusive shot_spec or asset_spec."""
    manifest, _, _ = compile_job(_mixed_sheet(tmp_path))
    p2 = build_execution_plan(manifest, pipeline="p2")

    for step in p2["steps"]:
        inputs = step["inputs"]
        unit_id = inputs["shot_id"]
        unit_kind = inputs.get("unit_kind")

        if step["tool"] == "pdx.media.technical_conform":
            assert unit_id in {"task_001", "task_003"}
            assert unit_kind == "shot"
            continue

        if unit_id == "task_002":
            # Asset unit (studio.asset.3d, studio.motion.mocap)
            assert unit_kind == "asset"
            assert "asset_spec" in inputs
            assert inputs["asset_spec"]["schema_version"] == "fsf_asset_spec_v1"
            assert inputs["asset_spec"]["kind"] == "unspecified"
            assert "shot_spec" not in inputs
        else:
            # Shot units (task_001, task_003)
            assert unit_kind == "shot"
            assert "shot_spec" in inputs
            assert inputs["shot_spec"]["schema_version"] == "fsf_shot_spec_v1"
            assert "asset_spec" not in inputs


def test_fsf_run_excludes_asset_from_pending_on_resume(tmp_path: Path, capsys) -> None:
    """P1-2: Asset must not be counted as pending or remain permanently pending across resumes."""
    sheet = _mixed_sheet(tmp_path)
    out = tmp_path / "run_resume_test"

    manifest, _, _ = compile_job(sheet, out_dir=out)
    assert manifest["task_count"] == 3
    assert manifest["shot_count"] == 2
    assert manifest["asset_count"] == 1

    # First inspection before any run
    status_0 = inspect_existing_artifacts(manifest, out)
    assert status_0["task_002"]["resume_kind"] == "not_applicable"
    assert status_0["task_002"]["completed"] is False

    resumed_0, pending_0, completed_0 = build_resumed_job(manifest, out)
    assert resumed_0["pending_count"] == 2  # Only shots
    assert resumed_0["completed_count"] == 0
    assert resumed_0["skipped_count"] == 1
    assert len(pending_0) == 2
    assert [t["id"] for t in pending_0] == ["task_001", "task_003"]
    assert len(completed_0) == 0
    assert resumed_0["is_resumed"] is False

    # First run: generates 2 shots
    exit_1 = main(["run", str(sheet), "--out", str(out)])
    assert exit_1 == 0
    captured_1 = capsys.readouterr().out
    assert "Pending: 2 | Reused: 0 | Skipped: 1" in captured_1

    # Second inspection: shots are completed; asset is not_applicable
    status_1 = inspect_existing_artifacts(manifest, out)
    assert status_1["task_001"]["completed"] is True
    assert status_1["task_003"]["completed"] is True
    assert status_1["task_002"]["resume_kind"] == "not_applicable"
    assert status_1["task_002"]["completed"] is False

    resumed_1, pending_1, completed_1 = build_resumed_job(manifest, out)
    assert resumed_1["pending_count"] == 0  # CRITICAL: not 1!
    assert resumed_1["completed_count"] == 2
    assert resumed_1["skipped_count"] == 1
    assert len(pending_1) == 0
    assert len(completed_1) == 2
    assert resumed_1["is_resumed"] is True

    # Second run (resume): 0 pending, 2 reused, 1 skipped
    exit_2 = main(["run", str(sheet), "--out", str(out)])
    assert exit_2 == 0
    captured_2 = capsys.readouterr().out
    assert "Pending: 0 | Reused: 2 | Skipped: 1" in captured_2
    assert "Finished." in captured_2


def test_media_conformance_check_step_binds_compiled_expected(tmp_path: Path) -> None:
    """P1-3: pdx.media.technical_conform step must bind the compiled expected contract and digest."""
    manifest, _, _ = compile_job(_mixed_sheet(tmp_path))

    # Give task_001 custom non-duration expected contract parameters
    custom_expected = {
        "containers": ["mov", "mp4"],
        "video_codecs": ["prores", "h264"],
        "width": 1920,
        "height": 1080,
        "fps": 30,
        "fps_tolerance": 0.005,
        "duration_seconds": 7.5,
        "duration_tolerance_seconds": 0.1,
        "audio": "required",
        "max_black_frame_ratio": 0.0,
        "max_freeze_duration_seconds": 0.2,
    }
    manifest["tasks"][0]["expected_media_contract"] = custom_expected
    # Explicit operator contract — not an intake default.
    manifest["tasks"][0]["expected_media_contract_source"] = "explicit"
    expected_digest = canonical_contract_digest(custom_expected)

    plan = build_execution_plan(manifest, executor_tool="studio.video.generate")
    qc_step = next(
        step for step in plan["steps"]
        if step["tool"] == "pdx.media.technical_conform" and step["inputs"]["shot_id"] == "task_001"
    )

    qc_inputs = qc_step["inputs"]
    assert qc_inputs["shot_id"] == "task_001"
    assert qc_inputs["duration_seconds"] == 7.5
    assert qc_inputs["expected_media_contract_digest"] == expected_digest

    bound_contract = qc_inputs["expected_media_contract"]
    # Non-duration fields must not drift back to defaults (1024, 576, 24, forbidden)
    assert bound_contract["width"] == 1920
    assert bound_contract["height"] == 1080
    assert bound_contract["fps"] == 30
    assert bound_contract["fps_tolerance"] == 0.005
    assert bound_contract["audio"] == "required"
    assert bound_contract["containers"] == ["mov", "mp4"]
    assert bound_contract["video_codecs"] == ["prores", "h264"]
    assert bound_contract["max_black_frame_ratio"] == 0.0
    assert bound_contract["max_freeze_duration_seconds"] == 0.2
    assert canonical_contract_digest(bound_contract) == expected_digest
