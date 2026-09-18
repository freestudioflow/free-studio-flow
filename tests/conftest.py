"""Pytest test fixtures for Free Studio Flow."""

from __future__ import annotations

from pathlib import Path
import pytest
from docx import Document

PROMPTS = [
    "Locked-off wide shot of an overcast ocean. Grey water, small waves, a flat horizon, no land, no ships. Soft rain on the surface. Contemporary digital capture, cool grey, no vintage film. 5 seconds",
    "Locked-off medium shot inside a dark industrial plant. Thick white steam drifts from overhead pipes. A few warm practical lamps stay in place. Dust hangs in the air. No people, no control-panel text. Contemporary digital capture. 10.0 seconds",
    "Locked-off shot of a distant colorful nebula in deep space. Slow drift of gas clouds. Stars behind it stay in place. No ships, no people, no text. Contemporary digital capture. 5s",
]


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Create a standard prompt sheet docx file."""
    docx_path = tmp_path / "prompt_sheet.docx"
    doc = Document()
    for index, prompt in enumerate(PROMPTS):
        if index > 0:
            doc.add_paragraph("-----")
        doc.add_paragraph(prompt)
    doc.save(docx_path)
    return docx_path
