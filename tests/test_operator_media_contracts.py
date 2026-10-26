"""A starting media preset must not bind every model to its resolution."""
import pytest
from docx import Document
from fsf.intake.compiler import compile_job
from fsf.runtime.plan import build_execution_plan
from fsf.conformance.media import build_expected_contract


def sheet(tmp_path):
    doc=Document();doc.add_paragraph('SHOT: A sphere moves, 1 seconds')
    doc.add_paragraph('ID: SC01');doc.add_paragraph('REVIEW_REQUIRED: false')
    path=tmp_path/'sheet.docx';doc.save(path);return path


def test_explicit_contracts_bind_intake_and_plan_without_changing_defaults(tmp_path):
    path=sheet(tmp_path)
    old,_,_=compile_job(path)
    contract=build_expected_contract({'duration_seconds':1},width=256,height=144,fps=24)
    job,_,_=compile_job(path,media_contracts={'SC01':contract})
    assert old['tasks'][0]['expected_media_contract']['width']==1024
    assert job['tasks'][0]['expected_media_contract']==contract
    assert job['source_sha256']==old['source_sha256']
    plan=build_execution_plan(job,shape=False)
    qc=next(s for s in plan['steps'] if s['tool']=='pdx.media.technical_conform')
    assert qc['inputs']['expected_media_contract']==contract
    assert compile_job(path)[0]==old


def test_unknown_shot_and_invalid_contract_refused_before_any_execution(tmp_path):
    path=sheet(tmp_path)
    with pytest.raises(ValueError,match='Invalid operator media contracts'):
        compile_job(path,media_contracts={'missing':{}})
    with pytest.raises(ValueError,match='Invalid operator media contract'):
        compile_job(path,media_contracts={'SC01':{'width':256}})
