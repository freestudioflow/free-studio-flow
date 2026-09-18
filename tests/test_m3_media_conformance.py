"""Tests for FSF-M3: Media Conformance Consumption & Template Conformance."""

from __future__ import annotations

from pathlib import Path
from docx import Document

from pdx_adapter_media import ProbeUnavailable
from fsf.conformance.media import (
    build_expected_contract,
    evaluate_media_file,
    map_conformance_status_to_host,
)
from fsf.conformance.template import evaluate_table_template_conformance


class MockGoodProbe:
    def probe(self, path: Path) -> dict:
        return {
            "format": {"duration": "5.0", "format_name": "mov,mp4"},
            "streams": [
                {
                    "index": 0,
                    "codec_type": "video",
                    "codec_name": "h264",
                    "width": 1024,
                    "height": 576,
                    "avg_frame_rate": "24/1",
                    "nb_frames": "120",
                }
            ],
            "black_frame_ratio": 0.0,
            "freeze_ranges": [],
        }


class MockMismatchProbe:
    def probe(self, path: Path) -> dict:
        return {
            "format": {"duration": "5.0", "format_name": "mov,mp4"},
            "streams": [
                {
                    "index": 0,
                    "codec_type": "video",
                    "codec_name": "h264",
                    "width": 1920,
                    "height": 1080,
                    "avg_frame_rate": "24/1",
                    "nb_frames": "120",
                }
            ],
            "black_frame_ratio": 0.0,
            "freeze_ranges": [],
        }


class MockFailingProbe:
    def probe(self, path: Path) -> dict:
        raise ProbeUnavailable("ffprobe is not installed on this system")


def test_build_expected_contract() -> None:
    task = {"duration_seconds": 10.0}
    expected = build_expected_contract(task, width=1024, height=576, fps=24)
    assert expected["width"] == 1024
    assert expected["height"] == 576
    assert expected["duration_seconds"] == 10.0
    assert expected["audio"] == "forbidden"


def test_media_conformance_four_statuses(tmp_path: Path) -> None:
    """M3 Exit Criterion: Media technical conformance evaluation without private ffprobe reimplementation."""
    sample_media = tmp_path / "sample.mp4"
    sample_media.write_bytes(b"TEST_MEDIA_BYTES")

    expected = build_expected_contract({"duration_seconds": 5.0})

    # 1. Conforms
    rep_good, act_good = evaluate_media_file(sample_media, expected, probe_runner=MockGoodProbe())
    assert rep_good["evaluation_status"] == "conforms"
    assert act_good["host_action"] == "permit_next_step"
    assert act_good["create_binding"] is True

    # 2. Does Not Conform (Dimensions Mismatch)
    rep_bad, act_bad = evaluate_media_file(sample_media, expected, probe_runner=MockMismatchProbe())
    assert rep_bad["evaluation_status"] == "does_not_conform"
    assert any(i["code"] == "MEDIA_DIMENSIONS_MISMATCH" for i in rep_bad["issues"])
    assert act_bad["host_action"] == "block_next_product_step"

    # 3. Evaluation Failed (Analyzer fails)
    rep_fail, act_fail = evaluate_media_file(sample_media, expected, probe_runner=MockFailingProbe())
    assert rep_fail["evaluation_status"] == "evaluation_failed"
    assert act_fail["host_action"] == "fail_block_or_retry"
    assert act_fail["create_binding"] is True

    # 4. Not Evaluated (Explicit skip / no analyzer requested)
    rep_skip, act_skip = evaluate_media_file(
        sample_media, expected, probe_runner=MockGoodProbe(), skip_evaluation=True
    )
    assert rep_skip["evaluation_status"] == "not_evaluated"
    assert "request_id" not in rep_skip
    assert act_skip["host_action"] == "continue_without_technical_check"
    assert act_skip["create_binding"] is False


def test_template_conformance_check() -> None:
    """M3 Template Conformance: Reference vs Candidate table check."""
    doc_ref = Document()
    t_ref = doc_ref.add_table(rows=2, cols=2)
    t_ref.cell(0, 0).text = "Shot"
    t_ref.cell(0, 1).text = "Duration"
    t_ref.cell(1, 0).text = "task_001"
    t_ref.cell(1, 1).text = "5.0"

    doc_cand = Document()
    t_cand = doc_cand.add_table(rows=2, cols=2)
    t_cand.cell(0, 0).text = "Shot"
    t_cand.cell(0, 1).text = "Duration"
    t_cand.cell(1, 0).text = "task_001"
    t_cand.cell(1, 1).text = "5.0"

    from io import BytesIO
    b_ref = BytesIO()
    doc_ref.save(b_ref)
    b_cand = BytesIO()
    doc_cand.save(b_cand)

    res = evaluate_table_template_conformance(b_ref.getvalue(), b_cand.getvalue())
    assert res["conforms"] is True
    assert not res.get("issues")
