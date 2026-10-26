"""Remote/Handoff identity binding with real media and the existing ingest gate."""
import copy
import pytest
from fsf.runtime.handoff import build_workorder,ingest_workorder_result
from fsf.production.remote_handoff import definition_from_workorder,fulfill_remote_video
from fsf.production.contracts import ProductionError,seal,digest,publish
from fsf.core.edit_store import read_json
from test_production_control import movie


def work():
    return build_workorder(job_id='film',run_id='run',step_id='generate',attempt_id=1,
        tool='studio.video.generate',plan_digest='a'*64,inputs_digest='b'*64,
        target_artifact='artifacts/shot.mp4',task_payload={'prompt':'a sphere moves'})


def verified_dispatch(folder):
    publish(folder/'ledger.json',seal({'status':'verified','request_digest':'c'*64,'profile_digest':digest({}),
        'result_digest':'d'*64},'ledger_digest'))


def test_actual_media_projection_does_not_write_journal_or_approve_review(tmp_path,monkeypatch):
    emitted=work();definition=definition_from_workorder(emitted,parameters={},inputs={})
    request={**definition,'request_digest':'c'*64,'recipe':{'steps':[{'outputs':[{'artifact_id':'video','path':'movie.mp4','role':'video'}]}]}}
    dispatch=tmp_path/'dispatch';(dispatch/'accepted/outputs').mkdir(parents=True)
    verified_dispatch(dispatch)
    source=movie(dispatch/'accepted/outputs/movie.mp4',count=24)
    result={'status':'completed','result_digest':'d'*64}
    # This is a bridge-only test. Real native collection is tested separately;
    # the test never claims the fixture video came from a neural model.
    monkeypatch.setattr('fsf.remote.v2.check_authorization',lambda *a,**kw:request)
    monkeypatch.setattr('fsf.remote.v2.collect_result',lambda *a,**kw:result)
    output=tmp_path/'production';output.mkdir()
    value=fulfill_remote_video(emitted,request,profile={},store=None,dispatch_directory=dispatch,output_root=output)
    assert not value['production_journal_written'] and value['motion_review']=='not_reviewed'
    path=output/'handoff/results'/(emitted['workorder_id']+'.result.json')
    assert ingest_workorder_result(path,emitted,output)['status']=='completed'
    assert (output/'artifacts/shot.mp4').read_bytes()==source.read_bytes()
    assert not list(output.glob('*journal*')) and not list(output.rglob('*review*'))
    assert fulfill_remote_video(emitted,request,profile={},store=None,dispatch_directory=dispatch,output_root=output)==value
    (output/'artifacts/shot.mp4').write_bytes(b'changed')
    with pytest.raises(ProductionError):
        fulfill_remote_video(emitted,request,profile={},store=None,dispatch_directory=dispatch,output_root=output)


@pytest.mark.parametrize('field,value',[('job_id','another'),('attempt_id',2),('inputs_digest','e'*64)])
def test_unrelated_remote_request_cannot_be_relabelled_as_current_workorder(tmp_path,monkeypatch,field,value):
    emitted=work();request=definition_from_workorder(emitted,parameters={},inputs={})
    request['parent'][field]=value
    monkeypatch.setattr('fsf.remote.v2.collect_result',lambda *a,**kw:pytest.fail('must reject before collection'))
    with pytest.raises(ProductionError,match='WORKORDER_MISMATCH'):
        fulfill_remote_video(emitted,request,profile={},store=None,dispatch_directory=tmp_path,output_root=tmp_path)


def test_tampered_workorder_cannot_supply_a_new_parent_identity():
    emitted=work();emitted['target_artifact']='../escape.mp4'
    with pytest.raises(ProductionError):definition_from_workorder(emitted,parameters={},inputs={})


def test_cli_definition_uses_emitted_identity_and_never_starts_provider(tmp_path):
    import json,subprocess,sys
    for name,value in [('work',work()),('parameters',{}),('inputs',{})]:
        (tmp_path/(name+'.json')).write_text(json.dumps(value))
    out=tmp_path/'definition.json'
    child=subprocess.run([sys.executable,'-m','fsf.cli','production','workorder-definition',str(tmp_path/'work.json'),
        '--parameters',str(tmp_path/'parameters.json'),'--inputs',str(tmp_path/'inputs.json'),'--out',str(out)],
        capture_output=True,text=True,timeout=30)
    assert child.returncode==0 and read_json(out)['parent']['job_id']=='film'
    assert read_json(out)['parent']['inputs_digest']=='b'*64
    assert not list(tmp_path.glob('*ledger*')) and not list(tmp_path.glob('*journal*'))


def test_bridge_then_real_production_qc_and_edit_import_without_forged_journal(tmp_path,monkeypatch):
    from docx import Document
    from test_post_production import cli
    from test_handoff import _write_test_mp4
    from fsf.runtime.handoff import list_pending_workorders
    from fsf.runtime.post_sources import import_edit_selection
    from fsf.core.editing import apply_edit_patch
    from test_edit_identity import patch
    root=tmp_path/'production';root.mkdir()
    doc=Document();doc.add_paragraph('SHOT: Moving sphere, 1 seconds');doc.add_paragraph('ID: SC01')
    doc.add_paragraph('REVIEW_REQUIRED: false')
    sheet=tmp_path/'sheet.docx';doc.save(sheet)
    cli('intake',sheet,'--out',root);cli('plan',sheet,'--media','--out',root/'initial-plan.json')
    cli('production','bind-remote-contracts',root/'initial-plan.json','--out',root/'plan.json')
    first=cli('execute',root/'plan.json','--out',root,'--handoff',check=False)
    assert first.returncode==10
    emitted=list_pending_workorders(root)[0]
    before=read_json(root/'fsf_execution_journal.json')
    definition=definition_from_workorder(emitted,parameters={},inputs={})
    request={**definition,'request_digest':'c'*64,'recipe':{'steps':[{'outputs':[{'artifact_id':'video','path':'movie.mp4','role':'video'}]}]}}
    dispatch=tmp_path/'remote';(dispatch/'accepted/outputs').mkdir(parents=True)
    verified_dispatch(dispatch)
    expected=emitted['expected_media_contract']
    _write_test_mp4(dispatch/'accepted/outputs/movie.mp4',expected['duration_seconds'],expected['width'],expected['height'])
    monkeypatch.setattr('fsf.remote.v2.check_authorization',lambda *a,**kw:request)
    monkeypatch.setattr('fsf.remote.v2.collect_result',lambda *a,**kw:{'status':'completed','result_digest':'d'*64,
        'consistency_acknowledgement':{'acknowledged_digest':emitted['requirements_digest'],
        'capability_gap':None,'ref_assets_used':[],'continuity_method':'prompt_only'}})
    fulfill_remote_video(emitted,request,profile={},store=None,dispatch_directory=dispatch,output_root=root)
    assert read_json(root/'fsf_execution_journal.json')==before
    cli('execute',root/'plan.json','--out',root,'--handoff')
    assert list(root.rglob('*.evidence.json'))
    plan=read_json(root/'plan.json');producer=next(s for s in plan['steps'] if s.get('tool')=='studio.video.generate')
    configs={'schema_version':'fsf_source_executions_v1','executions':[{'execution_id':'production_A','root':str(root),
        'job':'job_manifest.json','plan':'plan.json','journal':'fsf_execution_journal.json'}]}
    edit=import_edit_selection({'schema_version':'fsf_edit_import_v1','edit_id':'film','clock_rate':{'num':24,'den':1},
        'selections':[{'source_execution_id':'production_A','producer_step_id':producer['id'],'artifact_ref':producer['outputs'][0]}]},configs)
    length=edit['clips'][0]['out_frame'];original=edit['clips'][0]['clip_id']
    changed=apply_edit_patch(edit,patch(edit,{'op':'split','clip_id':original,'at_frame':length//2,'new_clip_ids':['left','right']}))
    changed=apply_edit_patch(changed,patch(changed,{'op':'reorder','clip_ids':['right','left']}))
    assert changed['revision']==3 and [c['clip_id'] for c in changed['clips']]==['right','left']
    assert read_json(root/'fsf_execution_journal.json')['steps'][producer['id']]['execution_kind']=='handoff'
    assert not list(root.rglob('*review*.json'))
