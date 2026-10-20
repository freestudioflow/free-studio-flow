"""Real production evidence, generation-free edits and hostile source mutations."""
import copy
from fractions import Fraction
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys

from docx import Document
import pytest

from fsf.assembly.fast_assembly import require_ffmpeg
from fsf.conformance.digest import canonical_contract_digest as digest
from fsf.core.editing import create_edit, apply_edit_patch, timing_anchor_digest
from fsf.runtime.post_plan import compile_post_plan
from fsf.runtime.post_sources import SourceVerifier, PostError, file_sha
from fsf.runtime.post_runtime import verify_post_receipt
from fsf.runtime.plan_executor import PlanExecutor
from test_edit_identity import patch, relative_track

ROOT = Path(__file__).resolve().parents[1]


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=True, indent=2), encoding="utf-8")
    return path


def cli(*args, env=None, check=True):
    environment = os.environ.copy()
    environment["PYTHONPATH"] = str(ROOT / "src")
    environment["PYTHONIOENCODING"] = "utf-8"
    environment.update(env or {})
    result = subprocess.run([sys.executable, "-m", "fsf.cli", *map(str, args)], cwd=ROOT,
                            env=environment, capture_output=True, text=True, encoding="utf-8", timeout=180)
    if check:
        assert result.returncode == 0, result.stdout + result.stderr
    return result


@pytest.fixture(scope="module")
def production(tmp_path_factory):
    root = tmp_path_factory.mktemp("post_production")
    ffmpeg = require_ffmpeg()
    counter = root / "generation_count.txt"
    script = root / "generator.py"
    script.write_text('''import subprocess,sys,pathlib
out,duration,shot,ffmpeg,counter=sys.argv[1:]
index=int(shot[-1]); count=pathlib.Path(counter)
count.write_text(str(int(count.read_text())+1) if count.exists() else "1")
gains=["rr=1:gg=0.1:bb=0.1","rr=0.1:gg=1:bb=0.1","rr=0.1:gg=0.1:bb=1"]
subprocess.run([ffmpeg,"-nostdin","-y","-f","lavfi","-i","testsrc2=size=1024x576:rate=24",
"-f","lavfi","-i",f"sine=frequency={300+index*150}:sample_rate=48000",
"-vf","colorchannelmixer="+gains[index-1],"-t",duration,"-c:v","libx264","-preset","ultrafast",
"-pix_fmt","yuv420p","-c:a","aac","-ac","2",out],check=True,stdin=subprocess.DEVNULL,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=40)
''', encoding="utf-8")
    doc = Document()
    for i in range(1, 4):
        if i > 1: doc.add_paragraph("-----")
        doc.add_paragraph("SHOT: Moving colored pattern, 2 seconds")
        doc.add_paragraph(f"ID: T0{i}")
        doc.add_paragraph("REVIEW_REQUIRED: true")
    sheet = root / "sheet.docx"; doc.save(sheet)
    cli("intake", sheet, "--out", root)
    cli("plan", sheet, "--media", "--out", root / "plan.json")
    config = write(root / "executors.json", {"executors": [{"executor_id": "local.post.test", "transport": "local",
        "tool_name": "studio.video.generate", "kind": "process", "expected_suffix": ".mp4",
        "command_args": [sys.executable, str(script), "{output}", "{duration}", "{shot_id}", ffmpeg, str(counter)]}]})
    cli("execute", root / "plan.json", "--out", root, env={"FSF_EXECUTORS_FILE": str(config)})
    plan = json.loads((root / "plan.json").read_text())
    for step in plan["steps"]:
        if step.get("tool") == "studio.video.generate":
            cli("review", "approve", root, step["id"], "--notes", "Fixture approval")
    assert counter.read_text() == "3"
    return root


@pytest.fixture
def bundle(production, tmp_path):
    root = tmp_path / "production"
    shutil.copytree(production, root)
    plan = json.loads((root / "plan.json").read_text())
    journal = json.loads((root / "fsf_execution_journal.json").read_text())
    sources, clips = [], []
    for i, step in enumerate([s for s in plan["steps"] if s.get("tool") == "studio.video.generate"], 1):
        ref = {"source_execution_id": "exec_A", "source_job_id": journal["job_id"], "source_run_id": journal["run_id"],
            "source_plan_digest": digest(plan), "source_journal_digest": digest(journal), "producer_step_id": step["id"],
            "artifact_ref": step["outputs"][0], "artifact_sha256": journal["steps"][step["id"]]["output_files"][0]["sha256"], "purpose": "video"}
        sources.append({"source_id": f"S0{i}", "shot_id": f"T0{i}", "label": f"Shot {i}", "duration_frames": 48, "source_ref": ref})
        clips.append({"clip_id": f"C0{i}", "source_id": f"S0{i}", "label": f"Clip {i}", "in_frame": 0, "out_frame": 48, "lineage": []})
    edit = create_edit(edit_id="film", clock_rate={"num": 24, "den": 1}, sources=sources, clips=clips)
    configs = {"schema_version": "fsf_source_executions_v1", "executions": [{"execution_id": "exec_A", "root": str(root),
        "job": "job_manifest.json", "plan": "plan.json", "journal": "fsf_execution_journal.json"}]}
    return root, edit, configs


def test_verified_sources_and_post_plan_shape(bundle):
    root, edit, configs = bundle
    resolved = SourceVerifier(edit, configs).verify()
    assert set(resolved) == {"S01", "S02", "S03"}
    plan = compile_post_plan(edit)
    assert [s["kind"] for s in plan["steps"]] == ["transform", "transform", "tool", "transform"]
    assert all(s.get("tool", "").startswith("pdx.media.") for s in plan["steps"] if s["kind"] == "tool")


def test_cli_three_edits_preview_resume_and_revocation(bundle, tmp_path):
    root, edit, configs = bundle
    before = {p.relative_to(root).as_posix(): file_sha(p) for p in root.rglob("*") if p.is_file()}
    track = {"schema_version": "fsf_edit_subtitle_track_v1", "track_id": "SUB", "clock_rate": edit["clock_rate"],
        "timebase": "clip_relative", "cues": [{"cue_id": "Q1", "clip_id": "C02", "start_frame": 0, "end_frame": 12, "text": "Signal detected."}]}
    config_file = write(tmp_path / "sources.json", configs)
    selection = write(tmp_path / "selection.json", {"schema_version": "fsf_edit_import_v1", "edit_id": "film", "clock_rate": edit["clock_rate"],
        "selections": [{"source_execution_id": s["source_ref"]["source_execution_id"], "producer_step_id": s["source_ref"]["producer_step_id"],
            "artifact_ref": s["source_ref"]["artifact_ref"], "label": f"Clip {i}"} for i, s in enumerate(edit["sources"], 1)]})
    edit_file = tmp_path / "edit.json"
    edit = json.loads(cli("edit", "import", selection, "--source-executions", config_file, "--out", edit_file).stdout)["result"]
    initial_patch = write(tmp_path / "initial-patch.json", patch(edit, {"op": "upsert_subtitle_track", "track": track}))
    edit = json.loads(cli("edit", "patch", edit_file, "--patch", initial_patch).stdout)["result"]
    prior_bundles = []
    for cycle in range(3):
        if cycle:
            ops = ([{"op": "trim", "clip_id": "C02", "in_frame": 12, "out_frame": 36},
                {"op": "reorder", "clip_ids": ["C03", "C02", "C01"]}] if cycle == 1 else
                [{"op": "duplicate", "clip_id": "C01", "new_clip_id": "C04"},
                 {"op": "upsert_cue", "track_id": "SUB", "cue": {"cue_id": "Q1", "clip_id": "C02", "start_frame": 0, "end_frame": 12, "text": "The echo returns."}}])
            patch_file = write(tmp_path / "patch.json", patch(edit, *ops))
            edit = json.loads(cli("edit", "patch", edit_file, "--patch", patch_file).stdout)["result"]
        plan_file = tmp_path / f"post{cycle}.json"
        cli("edit", "plan", edit_file, "--out", plan_file, "--width", "320", "--height", "180", "--subtitles", "sidecar")
        result = json.loads(cli("execute", plan_file, "--out", tmp_path / "post", "--source-executions", config_file).stdout)
        assert result["success"], result
        assert result["generation_dispatches"] == {"process": 0, "api": 0, "handoff": 0}
        assert result["post_transform_dispatches"] == 3
        assert result["media_process_counts"]["ffmpeg_calls"] >= 5
        folder = tmp_path / "post" / result["artifact_directory"]
        assert verify_post_receipt(folder)
        assert (folder / "captions.srt").is_file()
        prior_bundles.append((folder, file_sha(folder / "movie.mp4")))
        again = json.loads(cli("execute", plan_file, "--out", tmp_path / "post", "--source-executions", config_file).stdout)
        assert again["completed_count"] == 0 and again["reused_count"] == 4, again
        assert again["post_transform_dispatches"] == 0
        assert again["media_process_counts"]["ffmpeg_calls"] == 1  # Version query only.
        preview = json.loads(cli("edit", "preview", edit_file, "--source-executions", config_file, "--out", tmp_path / "previews",
            "--revision", edit["revision"], "--digest", edit["edit_digest"], "--clip-id", "C02", "--end-frame", "12").stdout)["result"]
        assert preview["manifest"]["kind"] == "preview" and preview["manifest"]["edit_digest"] == edit["edit_digest"]
        assert preview["manifest"]["source_out_frame"] - preview["manifest"]["source_in_frame"] == 12
        assert preview["manifest"]["timeline_end_frame"] - preview["manifest"]["timeline_start_frame"] == 12
    assert all(file_sha(folder / "movie.mp4") == sha for folder, sha in prior_bundles)
    after = {p.relative_to(root).as_posix(): file_sha(p) for p in root.rglob("*") if p.is_file()}
    assert before == after
    assert (root / "generation_count.txt").read_text() == "3"
    cli("review", "reject", root, edit["sources"][0]["source_ref"]["producer_step_id"], "--needs-rework")
    failed = cli("execute", plan_file, "--out", tmp_path / "post", "--source-executions", config_file, check=False)
    assert failed.returncode != 0 and json.loads(failed.stdout)["status"] == "blocked"
    assert all(file_sha(folder / "movie.mp4") == sha for folder, sha in prior_bundles)


@pytest.mark.parametrize("mutation", ["plan", "journal", "job", "missing", "sha", "shape", "producer", "declaration", "qc_missing", "qc_tamper", "evidence_missing", "review", "review_scope", "duration", "root", "purpose", "extra_producer", "requirements"])
def test_source_authority_mutations_rejected(bundle, mutation):
    root, edit, configs = bundle
    source = edit["sources"][0]
    ref = source["source_ref"]
    journal_path = root / "fsf_execution_journal.json"
    journal = json.loads(journal_path.read_text())
    plan_path = root / "plan.json"; plan = json.loads(plan_path.read_text())
    producer = next(s for s in plan["steps"] if s["id"] == ref["producer_step_id"])
    qc = next(s for s in plan["steps"] if s.get("inputs", {}).get("target_artifact") == ref["artifact_ref"])
    if mutation == "plan": plan["request_id"] += "changed"; write(plan_path, plan)
    elif mutation == "journal": journal["run_id"] = "changed"; write(journal_path, journal)
    elif mutation == "job":
        job = json.loads((root / "job_manifest.json").read_text()); job["source_sha256"] = "0" * 64; write(root / "job_manifest.json", job)
    elif mutation == "missing": (root / ref["artifact_ref"].split(":", 1)[1]).unlink()
    elif mutation == "sha": (root / ref["artifact_ref"].split(":", 1)[1]).write_bytes(b"wrong")
    elif mutation == "shape": journal["steps"][producer["id"]]["execution_kind"] = "shape"; write(journal_path, journal); ref["source_journal_digest"] = digest(journal)
    elif mutation == "producer": ref["producer_step_id"] = "missing"
    elif mutation == "declaration": ref["artifact_ref"] = "artifact:other.mp4"
    elif mutation == "qc_missing": del journal["steps"][qc["id"]]; write(journal_path, journal); ref["source_journal_digest"] = digest(journal)
    elif mutation == "qc_tamper": (root / qc["outputs"][0].split(":", 1)[1]).write_text("{}")
    elif mutation == "evidence_missing": (root / (qc["outputs"][0].split(":", 1)[1] + ".evidence.json")).unlink()
    elif mutation == "review":
        path = root / "fsf_consistency_reviews.json"; value = json.loads(path.read_text()); value["reviews"][producer["id"]]["verdict"] = "rejected"; write(path, value)
    elif mutation == "review_scope": journal["steps"][producer["id"]]["review_required"] = False; write(journal_path, journal); ref["source_journal_digest"] = digest(journal)
    elif mutation == "duration": source["duration_frames"] = 1000
    elif mutation == "root": configs["executions"][0]["execution_id"] = "other"
    elif mutation == "purpose": ref["purpose"] = "audio"
    elif mutation == "extra_producer": plan["steps"].append(copy.deepcopy(producer)); write(plan_path, plan); ref["source_plan_digest"] = digest(plan); journal["plan_digest"] = digest(plan); write(journal_path, journal); ref["source_journal_digest"] = digest(journal)
    elif mutation == "requirements": journal["steps"][producer["id"]]["requirements_digest"] = "0" * 64; write(journal_path, journal); ref["source_journal_digest"] = digest(journal)
    with pytest.raises(PostError): SourceVerifier(edit, configs).verify()


def _execute(edit, configs, output, **settings):
    result = PlanExecutor(output, source_executions=configs).execute_plan(compile_post_plan(edit, width=320, height=180, **settings))
    assert result["success"], result
    return Path(output) / result["artifact_directory"]


def _frame(video, seconds, destination):
    result = subprocess.run([require_ffmpeg(), "-nostdin", "-ss", str(seconds), "-i", str(video), "-vf", "scale=320:180",
        "-frames:v", "1", "-f", "rawvideo", "-pix_fmt", "rgb24", "pipe:1"], check=True, capture_output=True, timeout=30)
    assert len(result.stdout) == 320 * 180 * 3
    return result.stdout


def test_real_frame_trim_order_and_original_audio(bundle, tmp_path):
    import math, struct, wave
    root, edit, configs = bundle
    edit = apply_edit_patch(edit, patch(edit, {"op": "trim", "clip_id": "C02", "in_frame": 12, "out_frame": 36},
        {"op": "reorder", "clip_ids": ["C03", "C02", "C01"]}))
    folder = _execute(edit, configs, tmp_path / "post")
    video = folder / "movie.mp4"
    for time, dominant, frequency in [(0.25, 2, 750), (2.25, 1, 600), (3.25, 0, 450)]:
        image = _frame(video, time, tmp_path / "frame.png")
        means = [sum(image[channel::3]) / (len(image) / 3) for channel in range(3)]
        assert means[dominant] > 2 * max(value for i, value in enumerate(means) if i != dominant)
        audio = tmp_path / "audio.wav"
        subprocess.run([require_ffmpeg(), "-nostdin", "-y", "-ss", str(time), "-i", str(video), "-t", "0.25", "-ac", "1", "-ar", "8000", "-c:a", "pcm_s16le", str(audio)], check=True, capture_output=True, timeout=30)
        with wave.open(str(audio)) as stream: values = struct.unpack("<" + "h" * stream.getnframes(), stream.readframes(stream.getnframes()))
        amplitude = lambda freq: abs(sum(v * complex(math.cos(2 * math.pi * freq * i / 8000), math.sin(2 * math.pi * freq * i / 8000)) for i, v in enumerate(values))) / len(values)
        assert amplitude(frequency) > 100 and amplitude(frequency) > 10 * amplitude(1500)
    actual = _frame(video, 2.0, tmp_path / "actual.png")
    native = root / edit["sources"][1]["source_ref"]["artifact_ref"].split(":", 1)[1]
    expected = _frame(native, 0.5, tmp_path / "expected.png")
    wrong = _frame(native, 0, tmp_path / "wrong.png")
    difference = lambda other: sum(abs(a - b) for a, b in zip(actual, other)) / len(actual)
    assert difference(expected) < 8 and difference(wrong) > difference(expected) + 1


def test_rational_clock_with_different_source_rate(bundle, tmp_path):
    from fsf.conformance.media import SafeFfprobeRunner
    root, edit, configs = bundle
    clock = {"num": 24000, "den": 1001}
    for source in edit["sources"]: source["duration_frames"] = 47
    for clip in edit["clips"]: clip["out_frame"] = 47
    edit = create_edit(edit_id="rational", clock_rate=clock, sources=edit["sources"], clips=edit["clips"])
    folder = _execute(edit, configs, tmp_path / "post")
    raw = SafeFfprobeRunner().probe(folder / "movie.mp4")
    video = next(s for s in raw["streams"] if s["codec_type"] == "video")
    assert Fraction(video["avg_frame_rate"]) == Fraction(24000, 1001)
    assert int(video["nb_frames"]) == 141


def test_burned_subtitles_change_actual_pixels(bundle, tmp_path):
    root, edit, configs = bundle
    track = {"schema_version": "fsf_edit_subtitle_track_v1", "track_id": "SUB", "clock_rate": edit["clock_rate"], "timebase": "clip_relative",
        "cues": [{"cue_id": "Q1", "clip_id": "C01", "start_frame": 0, "end_frame": 24, "text": "ECHO STATION"}]}
    edit = apply_edit_patch(edit, patch(edit, {"op": "upsert_subtitle_track", "track": track}))
    plain = _execute(edit, configs, tmp_path / "plain", subtitles="sidecar")
    burned = _execute(edit, configs, tmp_path / "burned", subtitles="burned")
    left = _frame(plain / "movie.mp4", 0.25, tmp_path / "plain.png")
    right = _frame(burned / "movie.mp4", 0.25, tmp_path / "burned.png")
    assert sum(abs(a - b) for a, b in zip(left, right)) / len(left) > 0.1


def test_tampered_receipt_and_deleted_output_fail_integrity(bundle, tmp_path):
    root, edit, configs = bundle
    folder = _execute(edit, configs, tmp_path / "post")
    value = json.loads((folder / "receipt.json").read_text())
    value["total_frames"] += 1; write(folder / "receipt.json", value)
    assert not verify_post_receipt(folder)
    (folder / "movie.mp4").unlink()
    assert not verify_post_receipt(folder)


def test_no_resume_creates_new_attempt_without_overwrite(bundle, tmp_path):
    root, edit, configs = bundle
    plan = compile_post_plan(edit, width=320, height=180)
    first = PlanExecutor(tmp_path / "post", source_executions=configs).execute_plan(plan)
    second = PlanExecutor(tmp_path / "post", source_executions=configs, no_resume=True).execute_plan(plan)
    assert first["success"] and second["success"] and first["attempt_id"] != second["attempt_id"]
    assert verify_post_receipt(tmp_path / "post" / first["artifact_directory"]) and verify_post_receipt(tmp_path / "post" / second["artifact_directory"])


@pytest.mark.parametrize("attack", ["operation", "depends", "output", "settings", "tool"])
def test_post_plan_cannot_dispatch_arbitrary_operations(bundle, tmp_path, attack):
    root, edit, configs = bundle
    plan = compile_post_plan(edit)
    if attack == "operation": plan["steps"][0]["transform"] = "run_script"
    elif attack == "depends": plan["steps"][1]["depends_on"] = [edit["sources"][0]["source_ref"]["producer_step_id"]]
    elif attack == "output": plan["steps"][0]["outputs"] = ["artifact:../escape.mp4"]
    elif attack == "settings": plan["steps"][1]["inputs"]["post_context"]["width"] = 2
    else: plan["steps"][2]["tool"] = "studio.video.generate"
    result = PlanExecutor(tmp_path / "post", source_executions=configs).execute_plan(plan)
    assert not result["success"] and result["completed_count"] == 0
    assert (root / "generation_count.txt").read_text() == "3"


def test_trim_and_subtitle_repair_are_atomic(bundle):
    root, edit, configs = bundle
    track = {"schema_version": "fsf_edit_subtitle_track_v1", "track_id": "SUB", "clock_rate": edit["clock_rate"], "timebase": "clip_relative",
        "cues": [{"cue_id": "Q", "clip_id": "C01", "start_frame": 0, "end_frame": 12, "text": "Echo"}]}
    edit = apply_edit_patch(edit, patch(edit, {"op": "upsert_subtitle_track", "track": track}))
    original = copy.deepcopy(edit)
    with pytest.raises(ValueError): apply_edit_patch(edit, patch(edit, {"op": "remove", "clip_id": "C01"}))
    assert edit == original
    result = apply_edit_patch(edit, patch(edit, {"op": "split", "clip_id": "C01", "new_clip_ids": ["CA", "CB"], "at_frame": 24},
        {"op": "upsert_cue", "track_id": "SUB", "cue": {"cue_id": "Q", "clip_id": "CA", "start_frame": 0, "end_frame": 12, "text": "Echo"}}))
    assert result["subtitle_tracks"][0]["cues"][0]["clip_id"] == "CA"
    result = apply_edit_patch(result, patch(result, {"op": "remove_cue", "track_id": "SUB", "cue_id": "Q"}, {"op": "remove", "clip_id": "CA"}))
    assert not result["subtitle_tracks"][0]["cues"]


def test_absolute_anchor_stays_explicit(bundle):
    root, edit, configs = bundle
    track = {"schema_version": "fsf_edit_subtitle_track_v1", "track_id": "SUB", "clock_rate": edit["clock_rate"], "timebase": "timeline_absolute",
        "timing_anchor_digest": timing_anchor_digest(edit), "cues": [{"cue_id": "Q", "start_frame": 0, "end_frame": 12, "text": "Echo"}]}
    edit = apply_edit_patch(edit, patch(edit, {"op": "upsert_subtitle_track", "track": track}))
    with pytest.raises(ValueError): apply_edit_patch(edit, patch(edit, {"op": "trim", "clip_id": "C01", "in_frame": 0, "out_frame": 24}))
    result = apply_edit_patch(edit, patch(edit, {"op": "trim", "clip_id": "C01", "in_frame": 0, "out_frame": 24}, {"op": "confirm_timeline_anchor", "track_id": "SUB"}))
    assert result["subtitle_tracks"][0]["timing_anchor_digest"] == timing_anchor_digest(result)


def test_preview_stale_rejected_without_old_output(bundle, tmp_path):
    root, edit, configs = bundle
    source = write(tmp_path / "edit.json", edit); config = write(tmp_path / "sources.json", configs)
    output = tmp_path / "previews"
    failed = cli("edit", "preview", source, "--source-executions", config, "--out", output, "--revision", "2", "--digest", edit["edit_digest"], "--clip-id", "C01", check=False)
    assert failed.returncode == 4 and not output.exists() and "preview_path" not in failed.stdout


@pytest.mark.parametrize("case", ["M01", "M02", "M04", "M05", "M06", "M07", "M08", "M09", "M10", "M11", "M12", "M13", "M15", "M16", "M17", "M18"])
def test_mapping_matrix_runtime_cases(bundle, tmp_path, case):
    root, edit, configs = bundle
    journal_path = root / "fsf_execution_journal.json"
    journal = json.loads(journal_path.read_text())
    ref = edit["sources"][0]["source_ref"]
    state = journal["steps"][ref["producer_step_id"]]
    path = root / ref["artifact_ref"].split(":", 1)[1]
    if case in ("M01", "M02"):
        clips = [edit["clips"][0]]
        if case == "M02": clips.append({**edit["clips"][0], "clip_id": "C04", "lineage": ["C01"]})
        small = create_edit(edit_id="small", clock_rate=edit["clock_rate"], sources=[edit["sources"][0]], clips=clips)
        assert len(SourceVerifier(small, configs).verify()) == 1
        assert len(compile_post_plan(small)["steps"][0]["outputs"]) == 2
        return
    if case == "M12":
        reordered = apply_edit_patch(edit, patch(edit, {"op": "reorder", "clip_ids": ["C03", "C02", "C01"]}))
        assert timing_anchor_digest(reordered) != timing_anchor_digest(edit)
        assert SourceVerifier(reordered, configs).verify()
        return
    if case in ("M10", "M15"):
        plan = compile_post_plan(edit)
        if case == "M10": plan["steps"][0]["outputs"].pop()
        else: plan["steps"][1]["depends_on"] = [ref["producer_step_id"]]
        assert not PlanExecutor(tmp_path / "post", source_executions=configs).execute_plan(plan)["success"]
        return
    if case == "M11":
        with pytest.raises(ValueError):
            create_edit(edit_id="bad", clock_rate=edit["clock_rate"], sources=edit["sources"], clips=[edit["clips"][0]] * 2)
        return
    if case == "M04": shutil.copyfile(path, root / "same-bytes.mp4"); path.unlink()
    elif case == "M05": del journal["steps"][ref["producer_step_id"]]
    elif case == "M06": state["output_files"].append(copy.deepcopy(state["output_files"][0]))
    elif case == "M07": state["output_files"][0]["sha256"] = ""
    elif case == "M08":
        plan_path = root / "plan.json"; plan = json.loads(plan_path.read_text())
        producer = next(s for s in plan["steps"] if s["id"] == ref["producer_step_id"])
        producer["tool"] = "studio.voice.dub"; write(plan_path, plan)
        ref["source_plan_digest"] = digest(plan); journal["plan_digest"] = digest(plan)
    elif case == "M09": state["execution_kind"] = "mock"
    elif case == "M13": path.write_bytes(b"changed")
    elif case == "M16": journal["unrelated_field"] = "changed"
    elif case == "M17":
        alias = copy.deepcopy(configs["executions"][0]); alias["execution_id"] = "alias"
        configs["executions"].append(alias)
        second = edit["sources"][1]
        second["source_ref"] = {**copy.deepcopy(ref), "source_execution_id": "alias"}
    elif case == "M18": configs["executions"][0]["journal"] = "../escape.json"
    if case in ("M05", "M06", "M07", "M08", "M09"):
        write(journal_path, journal)
        for source in edit["sources"]: source["source_ref"]["source_journal_digest"] = digest(journal)
    elif case == "M16": write(journal_path, journal)
    with pytest.raises(PostError): SourceVerifier(edit, configs).verify()


def test_M03_same_names_in_independent_real_jobs(bundle, tmp_path):
    root, edit, configs = bundle
    other = tmp_path / "other_production"
    shutil.copytree(root, other)
    config = json.loads((other / "executors.json").read_text())
    args = config["executors"][0]["command_args"]
    args[1] = str(other / "generator.py"); args[-1] = str(other / "generation_count.txt")
    config_file = write(other / "executors.json", config)
    cli("execute", other / "plan.json", "--out", other, "--no-resume", "--job-id", "job_B", "--run-id", "run_B", env={"FSF_EXECUTORS_FILE": str(config_file)})
    journal = json.loads((other / "fsf_execution_journal.json").read_text())
    for step in json.loads((other / "plan.json").read_text())["steps"]:
        if step.get("tool") == "studio.video.generate": cli("review", "approve", other, step["id"])
    new = copy.deepcopy(edit["sources"][0])
    new["source_id"] = "S_OTHER"
    new["source_ref"].update(source_execution_id="exec_B", source_job_id="job_B", source_run_id="run_B", source_journal_digest=digest(journal))
    new["source_ref"]["artifact_sha256"] = journal["steps"][new["source_ref"]["producer_step_id"]]["output_files"][0]["sha256"]
    edit["sources"].append(new)
    edit["clips"].append({**edit["clips"][0], "clip_id": "C_OTHER", "source_id": "S_OTHER"})
    config_B = {**configs["executions"][0], "execution_id": "exec_B", "root": str(other)}
    configs["executions"].append(config_B)
    resolved = SourceVerifier(edit, configs).verify()
    assert resolved["S01"]["path"].name == resolved["S_OTHER"]["path"].name
    assert resolved["S01"]["path"] != resolved["S_OTHER"]["path"]


def test_review_revoked_after_resolution_blocks_publish(bundle, tmp_path, monkeypatch):
    from fsf.runtime.post_runtime import PostContext
    root, edit, configs = bundle
    real = PostContext.execute
    def revoke(self, step, output):
        result = real(self, step, output)
        if step["transform"] == "fsf.edit.resolve_sources":
            path = root / "fsf_consistency_reviews.json"; reviews = json.loads(path.read_text())
            reviews["reviews"][edit["sources"][0]["source_ref"]["producer_step_id"]]["verdict"] = "rejected"
            write(path, reviews)
        return result
    monkeypatch.setattr(PostContext, "execute", revoke)
    result = PlanExecutor(tmp_path / "post", source_executions=configs).execute_plan(compile_post_plan(edit))
    assert not result["success"] and not list((tmp_path / "post").rglob("delivery"))


def test_required_burn_failure_does_not_publish_sidecar(bundle, tmp_path, monkeypatch):
    import fsf.runtime.post_render as renderer
    root, edit, configs = bundle
    track = {"schema_version": "fsf_edit_subtitle_track_v1", "track_id": "SUB", "clock_rate": edit["clock_rate"], "timebase": "clip_relative",
        "cues": [{"cue_id": "Q", "clip_id": "C01", "start_frame": 0, "end_frame": 24, "text": "Echo"}]}
    edit = apply_edit_patch(edit, patch(edit, {"op": "upsert_subtitle_track", "track": track}))
    original = renderer.run_ffmpeg
    def fail(arguments, cwd):
        if any("subtitles=" in str(arg) for arg in arguments): raise PostError("POST_SUBTITLE_FILTER")
        return original(arguments, cwd)
    monkeypatch.setattr(renderer, "run_ffmpeg", fail)
    result = PlanExecutor(tmp_path / "post", source_executions=configs).execute_plan(compile_post_plan(edit, width=320, height=180, subtitles="burned"))
    assert not result["success"] and not list((tmp_path / "post").rglob("delivery"))


def test_cancel_before_bundle_rename_leaves_no_delivery(bundle, tmp_path, monkeypatch):
    import fsf.runtime.post_runtime as runtime
    root, edit, configs = bundle
    original = runtime.os.rename
    def cancel(source, destination):
        if Path(destination).name == "delivery": raise KeyboardInterrupt()
        return original(source, destination)
    monkeypatch.setattr(runtime.os, "rename", cancel)
    with pytest.raises(KeyboardInterrupt):
        PlanExecutor(tmp_path / "post", source_executions=configs).execute_plan(compile_post_plan(edit, width=320, height=180))
    assert not list((tmp_path / "post").rglob("delivery"))


def test_post_never_uses_generation_registry(bundle, tmp_path):
    from fsf.executors.registry import ExecutorRegistry
    class RefusingRegistry(ExecutorRegistry):
        def get(self, *args): raise AssertionError("Unexpected generation executor dispatch")
    root, edit, configs = bundle
    result = PlanExecutor(tmp_path / "post", registry=RefusingRegistry(), source_executions=configs).execute_plan(compile_post_plan(edit, width=320, height=180))
    assert result["success"] and (root / "generation_count.txt").read_text() == "3"


def test_journal_format_changes_keep_canonical_authority(bundle):
    root, edit, configs = bundle
    path = root / "fsf_execution_journal.json"
    value = json.loads(path.read_text())
    path.write_bytes(b"\xef\xbb\xbf" + json.dumps(value, separators=(",", ":")).replace("\n", "\r\n").encode())
    assert SourceVerifier(edit, configs).verify()


def test_post_source_qc_captures_remaster_evidence(production, tmp_path):
    from fsf.runtime.post_sources import import_edit_selection
    doc = Document(); doc.add_paragraph("SHOT: Moving colored pattern, 2 seconds"); doc.add_paragraph("ID: T01")
    doc.add_paragraph("REVIEW_REQUIRED: true")
    sheet = tmp_path / "sheet.docx"; doc.save(sheet)
    cli("intake", sheet, "--out", tmp_path)
    cli("plan", sheet, "--media", "--skills", "generate,remaster", "--post-sources", "--out", tmp_path / "plan.json")
    config = json.loads((production / "executors.json").read_text())
    config["executors"][0]["command_args"][-1] = str(tmp_path / "generation_count.txt")
    config["executors"].append({"executor_id": "local.remaster.test", "transport": "local", "tool_name": "studio.video.remaster",
        "kind": "process", "expected_suffix": ".mp4", "command_args": [sys.executable, "-c", "import shutil,sys;shutil.copyfile(sys.argv[1],sys.argv[2])", "{input}", "{output}"]})
    config_path = write(tmp_path / "executors.json", config)
    result, _ = json.JSONDecoder().raw_decode(cli("execute", tmp_path / "plan.json", "--out", tmp_path, env={"FSF_EXECUTORS_FILE": str(config_path)}).stdout)
    assert result["success"], result
    plan = json.loads((tmp_path / "plan.json").read_text())
    remaster = next(s for s in plan["steps"] if s.get("tool") == "studio.video.remaster")
    cli("review", "approve", tmp_path, remaster["id"])
    configs = {"schema_version": "fsf_source_executions_v1", "executions": [{"execution_id": "E", "root": str(tmp_path), "job": "job_manifest.json", "plan": "plan.json", "journal": "fsf_execution_journal.json"}]}
    edit = import_edit_selection({"schema_version": "fsf_edit_import_v1", "edit_id": "remastered", "clock_rate": {"num": 24, "den": 1},
        "selections": [{"source_execution_id": "E", "producer_step_id": remaster["id"], "artifact_ref": remaster["outputs"][0]}]}, configs)
    assert edit["sources"][0]["source_ref"]["producer_step_id"] == remaster["id"]
    assert SourceVerifier(edit, configs).verify()


@pytest.mark.parametrize("change", ["duplicate", "unknown_execution", "unknown_producer", "undeclared_output"])
def test_import_rejects_invalid_selection_without_creating_edit(bundle, tmp_path, change):
    root, edit, configs = bundle
    ref = edit["sources"][0]["source_ref"]
    selection = {"source_execution_id": ref["source_execution_id"], "producer_step_id": ref["producer_step_id"], "artifact_ref": ref["artifact_ref"]}
    definition = {"schema_version": "fsf_edit_import_v1", "edit_id": "film", "clock_rate": edit["clock_rate"], "selections": [selection]}
    if change == "duplicate": definition["selections"].append(copy.deepcopy(selection))
    elif change == "unknown_execution": selection["source_execution_id"] = "unknown"
    elif change == "unknown_producer": selection["producer_step_id"] = "unknown"
    else: selection["artifact_ref"] = "artifact:undeclared.mp4"
    result = cli("edit", "import", write(tmp_path / "selection.json", definition), "--source-executions", write(tmp_path / "config.json", configs), "--out", tmp_path / "edit.json", check=False)
    assert result.returncode != 0 and not (tmp_path / "edit.json").exists()
    assert "Traceback" not in result.stderr and str(root) not in result.stdout


def test_assembled_source_rechecks_upstream_required_reviews(bundle):
    from fsf.runtime.post_sources import import_edit_selection
    root, edit, configs = bundle
    cli("plan", root / "sheet.docx", "--media", "--assembly", "--out", root / "plan.json")
    config = json.loads((root / "executors.json").read_text())
    config["executors"][0]["command_args"][-1] = str(root / "generation_count.txt")
    config_path = write(root / "executors.json", config)
    cli("execute", root / "plan.json", "--out", root, "--no-resume", env={"FSF_EXECUTORS_FILE": str(config_path)}, check=False)
    plan = json.loads((root / "plan.json").read_text())
    producers = [s for s in plan["steps"] if s.get("tool") == "studio.video.generate"]
    for step in producers: cli("review", "approve", root, step["id"])
    cli("execute", root / "plan.json", "--out", root, env={"FSF_EXECUTORS_FILE": str(config_path)})
    assembled = next(s for s in plan["steps"] if s.get("transform") == "fsf.assembly.fast_concat")
    imported = import_edit_selection({"schema_version": "fsf_edit_import_v1", "edit_id": "assembled", "clock_rate": {"num": 24, "den": 1},
        "selections": [{"source_execution_id": "exec_A", "producer_step_id": assembled["id"], "artifact_ref": assembled["outputs"][0]}]}, configs)
    assert SourceVerifier(imported, configs).verify()
    cli("review", "reject", root, producers[0]["id"])
    with pytest.raises(PostError): SourceVerifier(imported, configs).verify()


@pytest.mark.parametrize("message", ["Glyph 0x1234 not found", "fontselect: failed to find any fallback with glyph 0x0"])
def test_successful_ffmpeg_with_missing_glyph_or_font_is_rejected(monkeypatch, message):
    import fsf.runtime.post_render as renderer
    from types import SimpleNamespace
    monkeypatch.setattr(renderer, "run_safe_subprocess", lambda *a, **k: SimpleNamespace(returncode=0, stderr=message, stdout=""))
    with pytest.raises(PostError, match="POST_SUBTITLE_FONT"):
        renderer.run_ffmpeg(["-version"], None)
